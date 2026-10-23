"""Excel export: one workbook per PDF, one sheet per statement, plus Validation and Metadata sheets.

Sheet layout (header row is found by the reader by looking for the ``Label`` cell, so a few
metadata lines above it are fine):

    A Row | B Level | C Row type | D Label | E Path | numeric value columns | raw text columns |
    underline-flag columns | Rule above

Nil cells (``－``) are left empty in the numeric columns; the ``Raw:`` columns keep exactly what was
printed so that nil, ``0`` and ``△0`` stay distinguishable.
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import Finding, Table

HEADER_ROW = 6
FIXED_HEADERS = ["Row", "Level", "Row type", "Label", "Path"]
LABEL_COL = 4  # column D
NUMBER_FORMAT = "#,##0;[Red]-#,##0"
PATH_SEP = " > "

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_SECTION_FILL = PatternFill("solid", fgColor="DDEBF7")
_FAIL_FILL = PatternFill("solid", fgColor="F8CBAD")
_OK_FILL = PatternFill("solid", fgColor="C6E0B4")
_ROUND_FILL = PatternFill("solid", fgColor="FFF2CC")
_TOP = Border(top=Side(style="thin"))
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def sheet_name(statement: str, taken: Sequence[str]) -> str:
    base = _INVALID_SHEET_CHARS.sub("", statement)[:31] or "Table"
    name, n = base, 2
    while name in taken:
        name = f"{base[:28]}_{n}"
        n += 1
    return name


def build_workbook(tables: Sequence[Table], findings: Sequence[Finding]) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    names: list[str] = []
    for table in tables:
        # a continuation page has no title: name its sheet after the page instead
        name = sheet_name(table.statement or f"Page {table.page_no + 1}", names)
        names.append(name)
        _write_table(wb.create_sheet(name), table)
    _write_validation(wb.create_sheet("Validation"), findings)
    _write_metadata(wb.create_sheet("Metadata"), tables)
    return wb


def write_workbook(tables: Sequence[Table], findings: Sequence[Finding], path: str | Path) -> Path:
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(tables, findings).save(out)
    return out


def workbook_bytes(tables: Sequence[Table], findings: Sequence[Finding]) -> bytes:
    """The workbook as bytes, for a download button."""
    buffer = BytesIO()
    build_workbook(tables, findings).save(buffer)
    return buffer.getvalue()


def _write_table(ws: Worksheet, table: Table) -> None:
    n_cols = len(table.headers)
    ws["A1"] = table.statement or f"Page {table.page_no + 1}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"], ws["B2"] = "Unit", table.unit or ""
    ws["A3"], ws["B3"] = "Source", f"{table.source} (page {table.page_no + 1})"
    ws["A4"], ws["B4"] = "Revision", str(table.meta.get("revision") or "")
    # The real statement name (blank for an untitled continuation page); A1 is only a visible title.
    ws["A5"], ws["B5"] = "Statement", table.statement or None

    num_start = len(FIXED_HEADERS) + 1
    raw_start = num_start + n_cols
    und_start = raw_start + n_cols
    rule_col = und_start + n_cols
    headers = (
        FIXED_HEADERS
        + list(table.headers)
        + [f"Raw: {h}" for h in table.headers]
        + [f"Underlined: {h}" for h in table.headers]
        + ["Rule above"]
    )
    for c, text in enumerate(headers, start=1):
        cell = ws.cell(row=HEADER_ROW, column=c, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    for i, row in enumerate(table.rows):
        r = HEADER_ROW + 1 + i
        ws.cell(row=r, column=1, value=row.index)
        ws.cell(row=r, column=2, value=row.level)
        ws.cell(row=r, column=3, value=row.row_type)
        label = ws.cell(row=r, column=LABEL_COL, value=row.label)
        label.alignment = Alignment(indent=row.level)
        ws.cell(row=r, column=5, value=PATH_SEP.join(row.path))
        for j, value in enumerate(row.values):
            num = ws.cell(row=r, column=num_start + j)
            if value is not None and value.number is not None:
                num.value = value.number
            num.number_format = NUMBER_FORMAT
            ws.cell(row=r, column=raw_start + j, value=value.raw if value else None)
            underlined = j < len(row.underlined) and row.underlined[j]
            ws.cell(row=r, column=und_start + j, value="yes" if underlined else None)
        ws.cell(row=r, column=rule_col, value="yes" if row.rule_above else None)
        _style_row(ws, r, row.row_type, rule_col)

    for c in range(1, rule_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 12
    ws.column_dimensions["D"].width = 38
    ws.column_dimensions["E"].width = 50
    for j in range(n_cols):
        ws.column_dimensions[get_column_letter(num_start + j)].width = 20
        ws.column_dimensions[get_column_letter(raw_start + j)].width = 14
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=num_start)


def _style_row(ws: Worksheet, r: int, row_type: str, last_col: int) -> None:
    for c in range(1, last_col + 1):
        cell = ws.cell(row=r, column=c)
        if row_type == "section":
            cell.fill = _SECTION_FILL
            cell.font = Font(bold=True)
        elif row_type == "total":
            cell.font = Font(bold=True)
            cell.border = _TOP


def _write_validation(ws: Worksheet, findings: Sequence[Finding]) -> None:
    headers = [
        "Check",
        "Status",
        "Table",
        "Column",
        "Row",
        "Expected",
        "Actual",
        "Diff",
        "Tolerance",
        "Detail",
    ]
    for c, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
    fills = {"PASS": _OK_FILL, "PASS_ROUNDING": _ROUND_FILL, "FAIL": _FAIL_FILL}
    for i, f in enumerate(findings, start=2):
        values = [
            f.check,
            f.status,
            f.table,
            f.column,
            f.row,
            f.expected,
            f.actual,
            f.diff,
            f.tolerance,
            f.detail,
        ]
        for c, v in enumerate(values, start=1):
            ws.cell(row=i, column=c, value=v)
        if f.status in fills:
            ws.cell(row=i, column=2).fill = fills[f.status]
    for c, width in enumerate([30, 16, 18, 30, 40, 14, 14, 8, 10, 60], start=1):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A2"


def _write_metadata(ws: Worksheet, tables: Sequence[Table]) -> None:
    ws.append(["Statement", "Unit", "Source", "Page", "Revision", "Rows", "Period ends"])
    for c in range(1, 8):
        ws.cell(row=1, column=c).font = Font(bold=True)
    for t in tables:
        ws.append(
            [
                t.statement,
                t.unit,
                t.source,
                t.page_no + 1,
                str(t.meta.get("revision") or ""),
                len(t.rows),
                ", ".join(p or "?" for p in t.periods),
            ]
        )
