"""Read a workbook written by ``export`` back into ``Table`` objects.

The validator and the evaluator work on what is *in the file*, so they also catch export bugs and
hand edits. Only the sheet layout is assumed, not the parser that produced it.
"""

from __future__ import annotations

from pathlib import Path
from typing import cast

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .export import FIXED_HEADERS, HEADER_ROW, LABEL_COL, PATH_SEP
from .hierarchy import build_paths
from .labels import detect_kind
from .models import Row, RowType, Table, Value
from .normalize import parse_value, period_end

_NON_TABLE_SHEETS = {"Validation", "Metadata"}
_ROW_TYPES = {"section", "item", "total", "net"}


def read_workbook(path: str | Path) -> list[Table]:
    wb = load_workbook(path, data_only=True)
    return [_read_sheet(ws) for ws in wb.worksheets if ws.title not in _NON_TABLE_SHEETS]


def _header_cells(ws: Worksheet) -> list[str]:
    return [str(c.value) if c.value is not None else "" for c in ws[HEADER_ROW]]


def _read_sheet(ws: Worksheet) -> Table:
    headers_all = _header_cells(ws)
    if headers_all[LABEL_COL - 1] != "Label":
        raise ValueError(f"sheet {ws.title!r}: no 'Label' header at row {HEADER_ROW}")
    num_start = len(FIXED_HEADERS)
    n_cols = next(i for i, h in enumerate(headers_all[num_start:]) if h.startswith("Raw: "))
    headers = headers_all[num_start : num_start + n_cols]
    raw_start = num_start + n_cols
    und_start = raw_start + n_cols
    rule_col = und_start + n_cols

    labels: list[str] = []
    levels: list[int] = []
    drafts: list[tuple[RowType, list[Value | None], list[bool], bool]] = []
    for cells in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        if all(c is None for c in cells):
            continue
        row_type = str(cells[2])
        if row_type not in _ROW_TYPES:
            raise ValueError(f"sheet {ws.title!r}: bad row type {row_type!r}")
        values: list[Value | None] = []
        for j in range(n_cols):
            number, raw = cells[num_start + j], cells[raw_start + j]
            parsed = parse_value(str(raw)) if raw is not None else None
            if number is None:
                values.append(parsed)  # nil / text / empty: decided by the raw text alone
            else:
                # The numeric cell is authoritative for the number; raw text keeps the printed form.
                values.append(
                    Value(
                        raw=str(raw) if raw is not None else str(number),
                        number=int(number),
                        kind="number",
                        negative_zero=bool(parsed and parsed.negative_zero),
                    )
                )
        labels.append(str(cells[LABEL_COL - 1] or ""))
        levels.append(int(cells[1] or 0))
        drafts.append(
            (
                cast(RowType, row_type),
                values,
                [cells[und_start + j] == "yes" for j in range(n_cols)],
                cells[rule_col] == "yes",
            )
        )

    paths = build_paths(labels, levels)
    rows = [
        Row(
            index=i,
            label=labels[i],
            level=levels[i],
            path=paths[i],
            values=values,
            row_type=row_type,
            rule_above=rule_above,
            underlined=underlined,
        )
        for i, (row_type, values, underlined, rule_above) in enumerate(drafts)
    ]
    meta_revision = ws["B4"].value
    if ws["A5"].value == "Statement":
        statement = str(ws["B5"].value or "")
    else:  # workbooks written before the Statement row existed
        statement = str(ws["A1"].value or ws.title)
    return Table(
        statement=statement,
        unit=ws["B2"].value or None,
        headers=headers,
        periods=[period_end(h) for h in headers],
        rows=rows,
        source=str(ws["B3"].value or ""),
        meta={"revision": meta_revision or None, "path_sep": PATH_SEP},
        kind=detect_kind(statement, labels),
    )
